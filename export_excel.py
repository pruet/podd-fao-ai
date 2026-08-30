import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel():
    json_path = os.path.join(os.path.dirname(__file__), "diseases.json")
    output_excel_path = os.path.join(os.path.dirname(__file__), "..", "Document", "List_of_Diseases_and_Guidelines.xlsx")
    output_excel_path = os.path.abspath(output_excel_path)
    
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)
        
    wb = Workbook()
    ws = wb.active
    ws.title = "Diseases & Guidelines"
    
    # Enable grid lines
    ws.views.sheetView[0].showGridLines = True
    
    # Header definitions
    headers = [
        "Animal Species\n(ຊະນິດສັດ)",
        "Disease Name (EN)",
        "Disease Name (Lao)\n(ຊື່ພະຍາດ ພາສາລາວ)",
        "Disease Name (Thai)\n(ชื่อโรค ภาษาไทย)",
        "Key Symptoms & Clinical Signs\n(ອາການສຳຄັນ / อาการสำคัญ)",
        "Basic Response - Lao\n(ຄຳແນະນຳເບື້ອງຕົ້ນສຳລັບສັດຕະວະແພດບ້ານ)",
        "Basic Response - Thai\n(คำแนะนำเบื้องต้นสำหรับสัตวแพทย์/เจ้าหน้าที่)",
        "Basic Response - English\n(Basic Field Guidelines for Local Vets)"
    ]
    
    ws.append(headers)
    
    # Styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    data_font = Font(name="Calibri", size=10)
    animal_font = Font(name="Calibri", size=11, bold=True, color="1F4E78")
    
    border_thin = Side(border_style="thin", color="D9D9D9")
    data_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    
    row_idx = 2
    for animal, diseases in data.items():
        for disease in diseases:
            name_en = disease.get("name", "")
            name_lo = disease.get("name_lo", "")
            name_th = disease.get("name_th", "")
            
            symptoms_list = disease.get("symptoms", [])
            symptoms_text = "\n".join(f"• {s}" for s in symptoms_list)
            
            basic_resp = disease.get("basic_response", {})
            resp_lo_list = basic_resp.get("lo", [])
            resp_th_list = basic_resp.get("th", [])
            resp_en_list = basic_resp.get("en", [])
            
            resp_lo_text = "\n".join(f"{i+1}. {r}" for i, r in enumerate(resp_lo_list))
            resp_th_text = "\n".join(f"{i+1}. {r}" for i, r in enumerate(resp_th_list))
            resp_en_text = "\n".join(f"{i+1}. {r}" for i, r in enumerate(resp_en_list))
            
            row_data = [
                animal,
                name_en,
                name_lo,
                name_th,
                symptoms_text,
                resp_lo_text,
                resp_th_text,
                resp_en_text
            ]
            
            ws.append(row_data)
            
            # Apply row styles
            fill_color = "F2F5F9" if row_idx % 2 == 0 else "FFFFFF"
            row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
            
            for col_idx in range(1, 9):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = animal_font if col_idx == 1 else data_font
                cell.fill = row_fill
                cell.border = data_border
                
                if col_idx in [1, 2, 3, 4]:
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                    
            row_idx += 1
            
    # Style Header Row
    ws.row_dimensions[1].height = 40
    for col_idx in range(1, 9):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = data_border

    # Set column widths
    col_widths = {
        1: 20, # Animal Species
        2: 24, # Disease Name EN
        3: 26, # Disease Name Lao
        4: 26, # Disease Name Thai
        5: 45, # Symptoms
        6: 50, # Response Lao
        7: 50, # Response Thai
        8: 50  # Response English
    }
    
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    # Freeze header row & apply filter
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{row_idx-1}"
    
    wb.save(output_excel_path)
    print(f"Successfully generated Excel file at: {output_excel_path}")

if __name__ == "__main__":
    generate_excel()
