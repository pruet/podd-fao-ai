import os
import re
import json
import openpyxl

def convert_excel_to_json():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    excel_path = os.path.join(base_dir, "..", "Document", "LAHIS_Symptom_List 16.9.2026_ART_SB_LOU_Lastest update_V4.xlsx")
    excel_path = os.path.abspath(excel_path)
    output_json_path = os.path.join(base_dir, "symptom.json")

    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Source Excel file not found at: {excel_path}")

    wb = openpyxl.load_workbook(excel_path, data_only=True)

    # Sheet name in workbook is 'All Symptoms (Lastest update)'
    sheet_name = "All Symptoms (Lastest update)"
    if sheet_name not in wb.sheetnames:
        for s in wb.sheetnames:
            if "all symptoms" in s.lower() and "update" in s.lower():
                sheet_name = s
                break

    s_main = wb[sheet_name]
    s_lao = wb["Symptom-Lao"] if "Symptom-Lao" in wb.sheetnames else None

    # Load clean Lao and English names from 'Symptom-Lao' sheet if available
    lao_map = {}
    if s_lao:
        for r in range(4, s_lao.max_row + 1):
            no = s_lao.cell(r, 1).value
            if no is not None:
                lao_map[no] = {
                    "category": s_lao.cell(r, 2).value.strip() if s_lao.cell(r, 2).value else "",
                    "name_en": s_lao.cell(r, 3).value.strip() if s_lao.cell(r, 3).value else "",
                    "name_lo": s_lao.cell(r, 4).value.strip() if s_lao.cell(r, 4).value else "",
                }

    species_keys = [
        (4, "Chicken"),
        (5, "Goose"),
        (6, "Duck"),
        (7, "Pig"),
        (8, "Cattle"),
        (9, "Buffalo"),
        (10, "Sheep and Goat"),
        (11, "Dog"),
        (12, "Cat")
    ]

    def split_en_lo(text):
        if not text:
            return "", ""
        # Search for first Lao unicode character
        match = re.search(r"[\u0e80-\u0eff]", text)
        if not match:
            return text.strip(), ""
        idx = match.start()
        en_part = text[:idx].strip()
        lo_part = text[idx:].strip()
        return en_part, lo_part

    symptoms_list = []
    for r in range(4, s_main.max_row + 1):
        no = s_main.cell(r, 1).value
        if no is None:
            continue

        lo_info = lao_map.get(no, {})
        raw_category = (s_main.cell(r, 2).value or "").strip()
        raw_symptom = (s_main.cell(r, 3).value or "").strip()

        clean_category = lo_info.get("category") or raw_category
        clean_symptom_en = lo_info.get("name_en") or raw_symptom
        clean_symptom_lo = lo_info.get("name_lo") or ""

        item = {
            "no": int(no),
            "category": clean_category,
            "symptom_name": clean_symptom_en,
            "symptom_name_en": clean_symptom_en,
            "symptom_name_lo": clean_symptom_lo,
            "raw_category": raw_category,
            "raw_symptom_name": raw_symptom,
            "species_guidance": {}
        }

        for col_idx, sp_name in species_keys:
            cell_val = s_main.cell(r, col_idx).value
            text = (cell_val or "").strip()
            en_part, lo_part = split_en_lo(text)

            item[sp_name] = text
            item["species_guidance"][sp_name] = {
                "text": text,
                "en": en_part,
                "lo": lo_part
            }

        symptoms_list.append(item)

    with open(output_json_path, "w", encoding="utf-8") as f:
        json.dump(symptoms_list, f, ensure_ascii=False, indent=2)

    print(f"Successfully converted {len(symptoms_list)} symptoms to {output_json_path}")
    return symptoms_list

if __name__ == "__main__":
    convert_excel_to_json()
